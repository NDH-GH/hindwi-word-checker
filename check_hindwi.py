import os
import sys
import time
import urllib.parse
import re
import html
import csv
import argparse
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
import openpyxl
import requests

# Fix console encoding on Windows
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

INPUT_EXCEL = "Validated word list.xlsx"
PROGRESS_CSV = "hindwi_validation_results.csv"
OUTPUT_EXCEL = "Validated_word_list_with_hindwi.xlsx"
SHEET_NAME = "Validated_words"
SAVE_INTERVAL = 200  # Save Excel & CSV every 200 words
NUM_WORKERS = 6      # 6 concurrent threads (~4-5 words per second)

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
    'Accept-Language': 'en-US,en;q=0.9,hi;q=0.8',
}

lock = threading.Lock()


def load_checkpoint():
    """Load already processed words from CSV to avoid redundant requests."""
    processed = {}
    if os.path.exists(PROGRESS_CSV):
        try:
            with open(PROGRESS_CSV, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    word = row.get('Word', '').strip()
                    if word:
                        processed[word] = row
            print(f"[CHECKPOINT] Loaded {len(processed)} previously checked words from '{PROGRESS_CSV}'.")
        except Exception as e:
            print(f"[CHECKPOINT WARNING] Error reading {PROGRESS_CSV}: {e}")
    return processed


def append_results_to_csv(new_rows):
    """Appends batch of results to the progress CSV."""
    with lock:
        file_exists = os.path.exists(PROGRESS_CSV) and os.path.getsize(PROGRESS_CSV) > 0
        fieldnames = ['Index', 'Word', 'Hindwi_Status', 'Column_D_Note', 'Card_Word', 'Meaning', 'Meaning_URL']
        with open(PROGRESS_CSV, 'a', encoding='utf-8', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            if not file_exists:
                writer.writeheader()
            for r in new_rows:
                writer.writerow(r)


def query_hindwi(session, word, retries=3):
    """Queries hindwidictionary.com and extracts the first card word & meaning."""
    encoded_word = urllib.parse.quote(word)
    url = f"https://www.hindwidictionary.com/search?keyword={encoded_word}"
    
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(url, timeout=12)
            if resp.status_code == 200:
                resp_html = resp.text
                
                # Target exclusively the first card
                card_match = re.search(
                    r'<div class=["\']WordRelatedLists saveWord listingCardBg overflowAuto["\']>(.*?)</div>',
                    resp_html,
                    re.DOTALL
                )
                if card_match:
                    card_html = card_match.group(1)
                    li_match = re.search(
                        r'<li>\s*([^<\n\r]+?)\s*<p class=["\']nobullet["\']>(.*?)</p>',
                        card_html,
                        re.DOTALL
                    )
                    if li_match:
                        card_word = li_match.group(1).strip()
                        raw_meaning = li_match.group(2).strip()
                        meaning = html.unescape(re.sub(r'\s+', ' ', raw_meaning))
                        
                        # Extract link
                        url_match = re.search(r'data-url=["\']([^"\']*)["\']', card_html)
                        meaning_url = ""
                        if url_match:
                            slug = url_match.group(1).strip()
                            meaning_url = f"https://www.hindwidictionary.com{slug}" if slug.startswith('/') else slug
                        
                        is_match = (word.strip() == card_word.strip())
                        status = "OKAY" if is_match else "MISMATCH"
                        
                        # Formulate Column D note
                        if is_match:
                            col_d_note = f"OKAY: {meaning}"
                        else:
                            col_d_note = f"Suggested: {card_word} | Meaning: {meaning}"
                            
                        return {
                            'Hindwi_Status': status,
                            'Column_D_Note': col_d_note,
                            'Card_Word': card_word,
                            'Meaning': meaning,
                            'Meaning_URL': meaning_url
                        }
                    else:
                        return {
                            'Hindwi_Status': 'PARSE_ERROR',
                            'Column_D_Note': 'Could not parse first card',
                            'Card_Word': '',
                            'Meaning': '',
                            'Meaning_URL': ''
                        }
                else:
                    return {
                        'Hindwi_Status': 'NOT_FOUND',
                        'Column_D_Note': 'Not found in Hindwi Dictionary',
                        'Card_Word': '',
                        'Meaning': '',
                        'Meaning_URL': ''
                    }
            elif resp.status_code == 429:
                time.sleep(3 * attempt)
            else:
                time.sleep(1)
        except Exception as e:
            if attempt == retries:
                return {
                    'Hindwi_Status': 'ERROR',
                    'Column_D_Note': f"Error: {e}",
                    'Card_Word': '',
                    'Meaning': str(e),
                    'Meaning_URL': ''
                }
            time.sleep(1.0)
            
    return {
        'Hindwi_Status': 'ERROR',
        'Column_D_Note': 'Max retries exceeded',
        'Card_Word': '',
        'Meaning': 'Max retries exceeded',
        'Meaning_URL': ''
    }


def export_excel(words_to_process, checkpoint):
    """Generates the enriched Excel file with Column D populated."""
    with lock:
        try:
            wb_full = openpyxl.load_workbook(INPUT_EXCEL)
            ws_full = wb_full[SHEET_NAME]
            
            # Headers
            ws_full.cell(row=1, column=4, value="Hindwi_Notes_and_Suggestions")
            ws_full.cell(row=1, column=5, value="Hindwi_Status")
            ws_full.cell(row=1, column=6, value="Suggested_First_Card_Word")
            ws_full.cell(row=1, column=7, value="Meaning")
            ws_full.cell(row=1, column=8, value="Meaning_URL")
            
            for r_idx, word in words_to_process:
                data = checkpoint.get(word)
                if data:
                    ws_full.cell(row=r_idx, column=4, value=data.get('Column_D_Note', ''))
                    ws_full.cell(row=r_idx, column=5, value=data.get('Hindwi_Status', ''))
                    ws_full.cell(row=r_idx, column=6, value=data.get('Card_Word', ''))
                    ws_full.cell(row=r_idx, column=7, value=data.get('Meaning', ''))
                    ws_full.cell(row=r_idx, column=8, value=data.get('Meaning_URL', ''))
                    
            wb_full.save(OUTPUT_EXCEL)
            print(f"[AUTO-SAVE] Enriched Excel saved to: '{OUTPUT_EXCEL}'")
        except Exception as e:
            print(f"Warning: Could not save Excel: {e}")


def worker_task(session, item):
    row_num, word = item
    result = query_hindwi(session, word)
    return {
        'Index': row_num,
        'Word': word,
        'Hindwi_Status': result['Hindwi_Status'],
        'Column_D_Note': result['Column_D_Note'],
        'Card_Word': result['Card_Word'],
        'Meaning': result['Meaning'],
        'Meaning_URL': result['Meaning_URL']
    }


def main():
    parser = argparse.ArgumentParser(description="Hindwi Dictionary Word Checker")
    parser.add_argument("--limit", type=int, default=0, help="Limit number of new words to process (0 = all)")
    parser.add_argument("--workers", type=int, default=NUM_WORKERS, help="Number of concurrent worker threads")
    args = parser.parse_args()

    if not os.path.exists(INPUT_EXCEL):
        print(f"Error: Could not find '{INPUT_EXCEL}' in current directory.")
        sys.exit(1)

    print(f"Opening '{INPUT_EXCEL}' (Sheet: '{SHEET_NAME}')...")
    wb = openpyxl.load_workbook(INPUT_EXCEL, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"Error: Sheet '{SHEET_NAME}' not found in {wb.sheetnames}")
        sys.exit(1)

    ws = wb[SHEET_NAME]
    
    # Read words from column 1
    words_to_process = []
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r_idx == 1:
            continue  # Skip header
        if not row or not row[0]:
            continue
        word = str(row[0]).strip()
        words_to_process.append((r_idx, word))

    total = len(words_to_process)
    checkpoint = load_checkpoint()

    # Filter unvisited words
    pending_items = [item for item in words_to_process if item[1] not in checkpoint]
    if args.limit > 0:
        pending_items = pending_items[:args.limit]

    print(f"Total words: {total} | Already completed: {len(checkpoint)} | Remaining to check: {len(pending_items)}")

    if not pending_items:
        print("All words already processed!")
        export_excel(words_to_process, checkpoint)
        return

    print(f"Launching {args.workers} concurrent workers...")

    # Create thread-local session
    thread_local = threading.local()

    def get_session():
        if not hasattr(thread_local, "session"):
            s = requests.Session()
            s.headers.update(HEADERS)
            thread_local.session = s
        return thread_local.session

    def process_item(item):
        session = get_session()
        return worker_task(session, item)

    batch_to_save = []
    completed_in_run = 0
    total_completed = len(checkpoint)

    t_start = time.time()

    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        futures = {executor.submit(process_item, item): item for item in pending_items}
        
        for future in as_completed(futures):
            record = future.result()
            word = record['Word']
            checkpoint[word] = record
            batch_to_save.append(record)
            
            completed_in_run += 1
            total_completed += 1
            
            # Print brief progress every 10 words
            if completed_in_run % 10 == 0 or completed_in_run == len(pending_items):
                elapsed = time.time() - t_start
                rate = completed_in_run / elapsed if elapsed > 0 else 0
                remaining_sec = (len(pending_items) - completed_in_run) / rate if rate > 0 else 0
                print(f"[{total_completed}/{total}] Checked: {word:12s} ({record['Hindwi_Status']:8s}) | Speed: {rate:.1f} w/s | Est. remaining: {remaining_sec/60:.1f} min")

            # Periodic batch save
            if len(batch_to_save) >= SAVE_INTERVAL:
                append_results_to_csv(batch_to_save)
                batch_to_save.clear()
                export_excel(words_to_process, checkpoint)

    if batch_to_save:
        append_results_to_csv(batch_to_save)
        batch_to_save.clear()

    # Final full Excel export
    export_excel(words_to_process, checkpoint)
    print(f"\n[COMPLETE] Successfully verified all words! Final file: '{OUTPUT_EXCEL}'")


if __name__ == '__main__':
    main()
